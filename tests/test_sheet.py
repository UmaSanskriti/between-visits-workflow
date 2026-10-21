"""Two-way Excel sync of practice information (data/practices.xlsx)."""

from __future__ import annotations

import json
from datetime import UTC, datetime

from conftest import CLEARANCES, SEED, SUPPRESSIONS
from openpyxl import load_workbook

from infra import sheet as sheetmod
from infra.brain import MockBrain
from infra.clock import FixedClock
from infra.config import load_config
from infra.email.enrich import SeedEnricher, enrich_contacts
from infra.pipeline import run_pipeline
from infra.sourcing import CsvSeedAdapter, SheetSeedAdapter
from infra.store import Store
from infra.voice import MockVoiceProvider

NOW = datetime(2026, 10, 20, 14, 0, tzinfo=UTC)
ROOT = SEED.parents[1]


def make_store(tmp_path):
    cfg = load_config(
        ROOT / "config.toml",
        env={
            "BV_DATA_DIR": str(tmp_path / "data"),
            "BV_EXPORT_DIR": str(tmp_path / "export"),
            "BV_OUTREACH_CHANNEL": "email",
        },
    )
    clock = FixedClock(NOW)
    store = Store(cfg.data_dir, clock).load()
    run_pipeline(
        store,
        cfg,
        clock,
        MockBrain(store, clock),
        MockVoiceProvider(cfg.webhook_secret, clock),
        SEED,
        CLEARANCES,
        SUPPRESSIONS,
        stages=("ingest", "verify"),
        trigger="test",
    )
    enrich_contacts(store, clock, [SeedEnricher(ROOT / "seeds" / "enrichment.csv", clock)])
    store.save()
    return cfg, clock, store


def cell(ws, header, col, row_idx):
    return ws.cell(row=row_idx, column=header.index(col) + 1)


def test_export_has_exact_columns_tags_and_no_outreach_state(tmp_path):
    cfg, clock, store = make_store(tmp_path)
    path = tmp_path / "practices.xlsx"
    n = sheetmod.export_sheet(store, path, clock)
    assert n == len(store.contacts) == 21
    ws = load_workbook(path)[sheetmod.SHEET_NAME]
    header = [c.value for c in ws[1]]
    assert header == sheetmod.COLUMNS
    assert not {"status", "attempts", "scenario", "call_clearance", "outcome"} & set(header)
    tags = {cell(ws, header, "source_tag", r).value for r in range(2, ws.max_row + 1)}
    assert tags == {"dummy"}
    patel = next(r for r in range(2, ws.max_row + 1) if cell(ws, header, "external_id", r).value == "t_01")
    assert cell(ws, header, "doctor_name", patel).value == "Priya Patel"
    assert cell(ws, header, "phone", patel).value == "+16175550101"
    assert cell(ws, header, "summary", patel).value  # enrichment travels with the practice
    assert ws.freeze_panes == "A2"
    side = json.loads(sheetmod.sidecar_path(path).read_text())
    assert side["fingerprint"] == sheetmod.fingerprint(path) and side["rows"] == 21
    assert not sheetmod.changed_since_export(path)


def test_round_trip_is_unchanged_and_edits_flow_back_with_audit(tmp_path):
    cfg, clock, store = make_store(tmp_path)
    path = tmp_path / "practices.xlsx"
    sheetmod.export_sheet(store, path, clock)
    counts = sheetmod.import_sheet(store, path, clock, log=lambda *a: None)
    assert counts["created"] == 0 and counts["updated"] == 0 and counts["unchanged"] == 21

    wb = load_workbook(path)
    ws = wb[sheetmod.SHEET_NAME]
    header = [c.value for c in ws[1]]
    patel = next(r for r in range(2, ws.max_row + 1) if cell(ws, header, "external_id", r).value == "t_01")
    cell(ws, header, "practice_name", patel).value = "Patel Family Medicine & Wellness"
    cell(ws, header, "summary", patel).value = "Edited in Excel: a solo practice with a new GLP-1 clinic."
    cell(ws, header, "hook", patel).value = "you opened a GLP-1 clinic this spring"
    cell(ws, header, "email", patel).value = "frontdesk@patelfamily.test"
    cell(ws, header, "glp1_mentioned", patel).value = "yes"
    wb.save(path)
    assert sheetmod.changed_since_export(path)

    counts = sheetmod.import_sheet(store, path, clock, log=lambda *a: None)
    assert counts["updated"] == 1
    contact = next(c for c in store.contacts.values() if c.external_id == "t_01")
    assert (
        contact.practice_name == "Patel Family Medicine & Wellness"
        and contact.hook == "you opened a GLP-1 clinic this spring"
    )
    en = store.enrichments[contact.id]
    assert en.source == "sheet" and en.summary.startswith("Edited in Excel") and en.glp1_mentioned is True
    emails = [ch.value for ch in store.channels_for(contact.id, "email")]
    assert "frontdesk@patelfamily.test" in emails
    new = next(ch for ch in store.channels_for(contact.id, "email") if ch.value == "frontdesk@patelfamily.test")
    assert new.verified is False and new.source == "sheet"  # must pass verify + clearance like any channel
    actions = [e["action"] for e in store.audit_log if e.get("actor") == "sheet"]
    assert "sheet_edit" in actions and "create" in actions and "import" in actions
    # a later seed/web enrichment refresh must not overwrite the sheet edit
    enrich_contacts(store, clock, [SeedEnricher(ROOT / "seeds" / "enrichment.csv", clock)], refresh=True)
    assert store.enrichments[contact.id].summary.startswith("Edited in Excel")


def test_new_row_creates_contact_with_generated_id_written_back(tmp_path):
    cfg, clock, store = make_store(tmp_path)
    path = tmp_path / "practices.xlsx"
    sheetmod.export_sheet(store, path, clock)
    wb = load_workbook(path)
    ws = wb[sheetmod.SHEET_NAME]
    header = [c.value for c in ws[1]]
    row = {col: "" for col in header}
    row.update(
        {
            "source_tag": "real",
            "doctor_name": "Nina Rao",
            "credentials": "MD",
            "specialty": "family medicine",
            "practice_name": "Rao Primary Care",
            "city": "Newton",
            "state": "MA",
            "phone": "617-555-0177",
            "website": "https://raoprimary.example",
            "services": "primary care; weight management",
        }
    )
    ws.append([row[col] for col in header])
    wb.save(path)
    counts = sheetmod.import_sheet(store, path, clock, log=lambda *a: None)
    assert counts["created"] == 1
    contact = next(c for c in store.contacts.values() if c.practice_name == "Rao Primary Care")
    assert contact.external_id.startswith("sheet:rao-primary-care") and contact.source_tag == "real"
    assert contact.status == "SOURCED" and store.channels_for(contact.id, "phone")[0].value == "+16175550177"
    ws = load_workbook(path)[sheetmod.SHEET_NAME]
    ids = [cell(ws, header, "external_id", r).value for r in range(2, ws.max_row + 1)]
    assert contact.external_id in ids  # written back so the next import matches the same contact
    assert store.enrichments[contact.id].services == ["primary care", "weight management"]


def test_missing_rows_are_reported_not_deleted(tmp_path):
    cfg, clock, store = make_store(tmp_path)
    path = tmp_path / "practices.xlsx"
    sheetmod.export_sheet(store, path, clock)
    wb = load_workbook(path)
    ws = wb[sheetmod.SHEET_NAME]
    ws.delete_rows(2)
    wb.save(path)
    before = len(store.contacts)
    counts = sheetmod.import_sheet(store, path, clock, log=lambda *a: None)
    assert counts["missing_from_sheet"] == 1 and len(store.contacts) == before


def test_sync_imports_an_edited_file_before_exporting(tmp_path):
    cfg, clock, store = make_store(tmp_path)
    path = tmp_path / "practices.xlsx"
    sheetmod.export_sheet(store, path, clock)
    wb = load_workbook(path)
    ws = wb[sheetmod.SHEET_NAME]
    header = [c.value for c in ws[1]]
    cell(ws, header, "city", 2).value = "Edited City"
    ext = cell(ws, header, "external_id", 2).value
    wb.save(path)
    result = sheetmod.sync_sheet(store, path, clock, log=lambda *a: None)
    assert result["imported"] and result["imported"]["updated"] == 1
    assert next(c for c in store.contacts.values() if c.external_id == ext).city == "Edited City"
    ws = load_workbook(path)[sheetmod.SHEET_NAME]
    assert cell(ws, header, "city", 2).value == "Edited City"  # the edit survived the export
    assert sheetmod.sync_sheet(store, path, clock, log=lambda *a: None)["imported"] is None  # nothing changed


def test_sheet_seed_adapter_matches_csv_rows_and_ingests_to_the_same_ids(tmp_path):
    cfg, clock, store = make_store(tmp_path)
    path = tmp_path / "practices.xlsx"
    sheetmod.export_sheet(store, path, clock)
    csv_rows = {r["external_id"]: r for r in CsvSeedAdapter(SEED).fetch()}
    sheet_rows = {r["external_id"]: r for r in SheetSeedAdapter(path).fetch()}
    assert set(sheet_rows) == set(csv_rows)
    for key in ("name", "practice_name", "specialty", "city", "state", "email"):
        assert sheet_rows["t_01"][key] == csv_rows["t_01"][key]
    assert sheet_rows["t_01"]["phone"] == "+16175550101"  # normalized on export
    assert set(csv_rows["t_01"]) <= set(sheet_rows["t_01"]) | {"line_type", "scenario"}
    assert CsvSeedAdapter(path).fetch() == SheetSeedAdapter(path).fetch()  # .xlsx seeds dispatch to the sheet
    # ingesting from the sheet maps onto the existing contacts (same ids), creating nothing new
    from infra.sourcing import ingest

    before = len(store.contacts)
    counts = ingest(store, SheetSeedAdapter(path), clock, cfg.default_time_zone)
    assert counts["contacts_created"] == 0 and len(store.contacts) == before
