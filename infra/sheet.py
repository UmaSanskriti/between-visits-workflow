"""Two-way sync between the practice database and an Excel workbook (data/practices.xlsx).

The sheet is the *practice information* database only: the columns we would scrape from the internet, with a
dummy/real tag. Outreach state (attempts, status, clearances) never goes in; the dashboard owns that.

Rules:
- export_sheet writes every contact (+ first phone, first email, enrichment) and records a fingerprint sidecar.
- import_sheet upserts practice fields, channels and enrichment (source "sheet") by external_id; rows missing
  from the sheet are reported, never deleted.
- sync_sheet applies the conflict rule: if the file changed since our last export, import first (the sheet wins
  for practice fields), then export the merged state. An edit is never clobbered.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import tomllib
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .clock import Clock, iso
from .models import Channel, Contact, ContactStatus, Enrichment
from .store import Store

ROOT = Path(__file__).resolve().parents[1]
SHEET_NAME = "practices"
COLUMNS = [
    "source_tag",
    "external_id",
    "npi",
    "doctor_name",
    "credentials",
    "specialty",
    "practice_name",
    "practice_size",
    "address",
    "city",
    "state",
    "time_zone",
    "phone",
    "email",
    "website",
    "hook",
    "summary",
    "services",
    "glp1_mentioned",
    "weight_management_mentioned",
    "doctor_bio",
    "patient_focus",
    "recent_note",
    "languages",
    "pronunciation_hint",
    "source",
    "source_url",
    "retrieved_at",
]
CONTACT_FIELDS = {  # sheet column -> Contact attribute
    "doctor_name": "full_name",
    "credentials": "credentials",
    "specialty": "specialty",
    "practice_name": "practice_name",
    "practice_size": "practice_size",
    "address": "address",
    "city": "city",
    "state": "state",
    "time_zone": "time_zone",
    "website": "website",
    "hook": "hook",
    "pronunciation_hint": "pronunciation_hint",
}
ENRICH_FIELDS = ["summary", "doctor_bio", "patient_focus", "recent_note", "languages", "website"]
WIDTHS = {
    "source_tag": 10,
    "external_id": 14,
    "npi": 12,
    "doctor_name": 20,
    "practice_name": 30,
    "address": 24,
    "phone": 15,
    "email": 26,
    "website": 28,
    "hook": 40,
    "summary": 50,
    "services": 36,
    "doctor_bio": 50,
    "patient_focus": 28,
    "recent_note": 36,
    "source_url": 40,
    "retrieved_at": 20,
}


def sheet_path(cfg=None, env: dict | None = None) -> Path:
    """BV_SHEET env, else [sheet].path in config.toml, else <data_dir>/practices.xlsx."""
    env = os.environ if env is None else env
    if env.get("BV_SHEET"):
        return Path(env["BV_SHEET"])
    config = ROOT / "config.toml"
    if config.exists():
        raw = tomllib.loads(config.read_text())
        configured = (raw.get("sheet") or {}).get("path")
        if configured:
            return Path(configured)
    data_dir = Path(getattr(cfg, "data_dir", "data")) if cfg is not None else Path("data")
    return data_dir / "practices.xlsx"


def sidecar_path(path: Path) -> Path:
    return path.with_name(path.stem + ".sheet.json")


def source_tag_for(contact: Contact) -> str:
    if contact.source_tag:
        return contact.source_tag
    if contact.lead_source.startswith("csv:test_doctors") or contact.external_id.startswith("t_"):
        return "dummy"
    if contact.lead_source in ("npi_registry", "web") or contact.external_id.startswith("npi:"):
        return "real"
    return "real" if contact.lead_source else "dummy"


def _yes(value) -> bool:
    return str(value or "").strip().lower() in ("yes", "y", "true", "1")


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")[:40] or "practice"


def fingerprint(path: Path) -> str:
    return hashlib.sha1(path.read_bytes()).hexdigest()[:16] if path.exists() else ""


def read_sidecar(path: Path) -> dict:
    sc = sidecar_path(path)
    return json.loads(sc.read_text()) if sc.exists() else {}


def write_sidecar(path: Path, **fields) -> dict:
    data = read_sidecar(path)
    data.update(fields)
    sidecar_path(path).write_text(json.dumps(data, indent=1, sort_keys=True) + "\n")
    return data


def changed_since_export(path: Path) -> bool:
    """True when the workbook differs from what we last wrote (someone edited it)."""
    if not path.exists():
        return False
    side = read_sidecar(path)
    return bool(side.get("fingerprint")) and fingerprint(path) != side["fingerprint"]


def row_for(store: Store, contact: Contact) -> dict:
    phones = [ch for ch in store.channels_for(contact.id, "phone")]
    emails = [ch for ch in store.channels_for(contact.id, "email")]
    en = store.enrichments.get(contact.id)
    phone = (phones[0].value or phones[0].raw_value) if phones else ""
    email = emails[0].value if emails else ""
    source = contact.lead_source
    source_url = phones[0].source_url if phones else (emails[0].source_url if emails else "")
    retrieved = phones[0].retrieved_at if phones else (emails[0].retrieved_at if emails else "")
    return {
        "source_tag": source_tag_for(contact),
        "external_id": contact.external_id,
        "npi": contact.external_id[4:] if contact.external_id.startswith("npi:") else "",
        "doctor_name": contact.full_name,
        "credentials": contact.credentials,
        "specialty": contact.specialty,
        "practice_name": contact.practice_name,
        "practice_size": contact.practice_size,
        "address": contact.address,
        "city": contact.city,
        "state": contact.state,
        "time_zone": contact.time_zone,
        "phone": phone,
        "email": email,
        "website": contact.website or (en.website if en else ""),
        "hook": contact.hook or (en.hook if en else ""),
        "summary": en.summary if en else "",
        "services": "; ".join(en.services) if en else "",
        "glp1_mentioned": ("yes" if en.glp1_mentioned else "no") if en else "",
        "weight_management_mentioned": ("yes" if en.weight_management_mentioned else "no") if en else "",
        "doctor_bio": en.doctor_bio if en else "",
        "patient_focus": en.patient_focus if en else "",
        "recent_note": en.recent_note if en else "",
        "languages": en.languages if en else "",
        "pronunciation_hint": contact.pronunciation_hint,
        "source": source,
        "source_url": source_url,
        "retrieved_at": retrieved,
    }


def export_sheet(store: Store, path: str | Path, clock: Clock | None = None) -> int:
    """Write the whole practice database to the workbook. Returns the row count."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(COLUMNS)
    ws.freeze_panes = "A2"
    contacts = sorted(store.contacts.values(), key=lambda c: (source_tag_for(c) != "dummy", c.external_id, c.full_name))
    for contact in contacts:
        row = row_for(store, contact)
        ws.append([row[col] for col in COLUMNS])
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 14)
    wb.save(path)
    now = iso((clock or Clock()).now())
    write_sidecar(
        path, last_export=now, export_mtime=path.stat().st_mtime, fingerprint=fingerprint(path), rows=len(contacts)
    )
    return len(contacts)


def read_sheet_rows(path: str | Path) -> list[dict]:
    """Rows as dicts keyed by COLUMNS (unknown extra columns are kept too)."""
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    rows: list[dict] = []
    header: list[str] = []
    for i, values in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            header = [_cell(v) for v in values]
            continue
        if not any(_cell(v) for v in values):
            continue
        row = {header[j]: _cell(values[j]) for j in range(min(len(header), len(values))) if header[j]}
        rows.append(row)
    wb.close()
    return rows


def _write_back_ids(path: Path, assignments: dict[int, str]) -> None:
    """Fill generated external_ids into rows that had none (row index -> id)."""
    if not assignments:
        return
    wb = load_workbook(path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    col = COLUMNS.index("external_id") + 1
    for row_idx, external_id in assignments.items():
        ws.cell(row=row_idx, column=col, value=external_id)
    wb.save(path)


def import_sheet(store: Store, path: str | Path, clock: Clock, log=print) -> dict:
    """Upsert practice information from the workbook (the sheet wins for these fields). Never deletes."""
    from .sourcing import contact_id_for  # local import: sourcing imports nothing from here
    from .verify import normalize_e164

    path = Path(path)
    now = iso(clock.now())
    counts = {"rows": 0, "created": 0, "updated": 0, "unchanged": 0, "missing_from_sheet": 0}
    by_external = {c.external_id: c for c in store.contacts.values() if c.external_id}
    seen: set[str] = set()
    generated: dict[int, str] = {}
    rows = read_sheet_rows(path)
    for idx, row in enumerate(rows, start=2):  # openpyxl row numbers (1 = header)
        counts["rows"] += 1
        external_id = row.get("external_id", "")
        if not external_id:
            external_id = "sheet:" + _slug(row.get("practice_name") or row.get("doctor_name") or f"row-{idx}")
            generated[idx] = external_id
            row["external_id"] = external_id
        seen.add(external_id)
        tag = (row.get("source_tag") or "").lower() or ("dummy" if external_id.startswith("t_") else "real")
        contact = by_external.get(external_id)
        changes: dict = {}
        if contact is None:
            seed_row = {
                "external_id": external_id,
                "name": row.get("doctor_name", ""),
                "practice_name": row.get("practice_name", ""),
                "phone": row.get("phone", ""),
                "email": row.get("email", ""),
            }
            contact = Contact(
                id=contact_id_for("csv_seed", seed_row),
                status=ContactStatus.SOURCED,
                external_id=external_id,
                lead_source=row.get("source") or "sheet",
                source_tag=tag,
                created_at=now,
                updated_at=now,
                **{attr: row.get(col, "") for col, attr in CONTACT_FIELDS.items()},
            )
            if not contact.time_zone:
                contact.time_zone = "America/New_York"
            store.add("contacts", contact, actor="sheet")
            by_external[external_id] = contact
            counts["created"] += 1
            created = True
        else:
            created = False
            en_existing = store.enrichments.get(contact.id)
            for col, attr in CONTACT_FIELDS.items():
                value = row.get(col, "")
                if col == "time_zone" and not value:
                    continue
                current = getattr(contact, attr)
                if col in ("hook", "website") and not current and en_existing:
                    current = getattr(en_existing, col)  # the sheet shows the effective value (enrichment fallback)
                if current != value:
                    changes[attr] = value
            if contact.source_tag != tag:
                changes["source_tag"] = tag
            if changes:
                store.update("contacts", contact, "sheet_edit", changes, actor="sheet")
        channel_changes = _upsert_channels(store, contact, row, normalize_e164, now)
        enrich_changes = _upsert_enrichment(store, contact, row, now)
        if created:
            continue
        if changes or channel_changes or enrich_changes:
            counts["updated"] += 1
        else:
            counts["unchanged"] += 1
    missing = [c.external_id for c in store.contacts.values() if c.external_id and c.external_id not in seen]
    counts["missing_from_sheet"] = len(missing)
    counts["missing_ids"] = missing[:50]
    if generated:
        _write_back_ids(path, generated)
    store.audit(
        "sheet", str(path), "import", {}, {k: v for k, v in counts.items() if k != "missing_ids"}, actor="sheet"
    )
    store.save()
    write_sidecar(
        path, last_import=now, fingerprint=fingerprint(path), export_mtime=path.stat().st_mtime, rows=len(rows)
    )
    log(
        f"sheet: imported {counts['rows']} rows ({counts['created']} new, {counts['updated']} updated, "
        f"{counts['missing_from_sheet']} in database but not in sheet)"
    )
    return counts


def _upsert_channels(store: Store, contact: Contact, row: dict, normalize_e164, now: str) -> bool:
    from .sourcing import short_hash

    changed = False
    for kind, raw in (("phone", row.get("phone", "")), ("email", row.get("email", ""))):
        if not raw:
            continue
        value = (normalize_e164(raw) or raw) if kind == "phone" else raw.lower()
        existing = store.channels_for(contact.id, kind)
        if any(ch.value == value or ch.raw_value == raw for ch in existing):
            continue
        if existing:  # the sheet replaced the number/address: keep history, add the new one unverified
            pass
        channel = Channel(
            id="ch_" + short_hash(contact.id, kind, raw),
            contact_id=contact.id,
            type=kind,
            value=value,
            raw_value=raw,
            source="sheet",
            source_url=row.get("source_url", "") or "practices.xlsx",
            retrieved_at=now,
        )
        store.add("channels", channel, actor="sheet", contact_id=contact.id)
        changed = True
    return changed


def _upsert_enrichment(store: Store, contact: Contact, row: dict, now: str) -> bool:
    has_any = any(
        row.get(col) for col in ENRICH_FIELDS + ["services", "glp1_mentioned", "weight_management_mentioned", "hook"]
    )
    if not has_any:
        return False
    services = [s.strip() for s in (row.get("services") or "").replace(",", ";").split(";") if s.strip()]
    fields = {
        "website": row.get("website", ""),
        "summary": row.get("summary", ""),
        "services": services,
        "glp1_mentioned": _yes(row.get("glp1_mentioned")),
        "weight_management_mentioned": _yes(row.get("weight_management_mentioned")),
        "doctor_bio": row.get("doctor_bio", ""),
        "patient_focus": row.get("patient_focus", ""),
        "recent_note": row.get("recent_note", ""),
        "languages": row.get("languages", ""),
        "hook": row.get("hook", ""),
    }
    existing = store.enrichments.get(contact.id)
    if existing is not None:
        for col in ("hook", "website"):  # exported as the effective value: unchanged when it still matches
            effective = getattr(contact, col) or getattr(existing, col)
            if fields[col] == effective:
                fields[col] = getattr(existing, col)
    if existing is None:
        store.add(
            "enrichments",
            Enrichment(
                id=contact.id,
                contact_id=contact.id,
                source="sheet",
                source_note="edited in practices.xlsx",
                retrieved_at=now,
                confidence="medium",
                created_at=now,
                updated_at=now,
                **fields,
            ),
            actor="sheet",
            contact_id=contact.id,
        )
        return True
    changes = {k: v for k, v in fields.items() if getattr(existing, k) != v}
    if changes:
        changes.update({"source": "sheet", "source_note": "edited in practices.xlsx", "retrieved_at": now})
        store.update("enrichments", existing, "sheet_edit", changes, actor="sheet", contact_id=contact.id)
        return True
    return False


def sync_sheet(store: Store, path: str | Path, clock: Clock, log=print) -> dict:
    """Conflict rule: import first if the file changed since our last export, then export the merged state."""
    path = Path(path)
    result = {"imported": None, "rows": 0}
    if changed_since_export(path):
        result["imported"] = import_sheet(store, path, clock, log)
    result["rows"] = export_sheet(store, path, clock)
    return result
